import io
import os
import re
import time
import streamlit as st
from google import genai
from google.genai import types

# Intentar importar python-docx para crear archivos .docx reales
try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# ==============================================================================
# ENLACES NORMATIVOS
# ==============================================================================
URL_CONSTITUCION_POLITICA = "https://www.registraduria.gov.co/IMG/pdf/constitucio-politica-colombia-1991.pdf"
URL_LEY_115 = (
    "https://www.mineducacion.gov.co/1621/articles-85906_archivo_pdf.pdf"
)
URL_LEY_1098 = (
    "https://www.icbf.gov.co/sites/default/files/codigoinfancialey1098.pdf"
)
URL_LEY_1620 = "https://www.funcionpublica.gov.co/eva/gestornormativo/norma_pdf.php?i=52287"
URL_DECRETO_1965 = "https://www.funcionpublica.gov.co/eva/gestornormativo/norma_pdf.php?i=54537"
URL_MANUAL_CONVIVENCIA = "https://drive.google.com/file/d/10WqGY5EvXzCMPROBZB6Ga6J4zjrEzmOG/view?usp=sharing"
# ==============================================================================

# Configuración inicial de la página
st.set_page_config(
    page_title="Sistema Integral de Convivencia Escolar - Institución Educativa Técnica Sagrado Corazón",
    layout="centered",
)

# Estilos CSS personalizados y optimizados
st.markdown(
    """
    <style>
    .stApp {
        background-color: #F3F4F6;
        color: #1F2937;
        font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
    }
    
    /* Barra lateral y todos sus textos y elementos en blanco */
    section[data-testid="stSidebar"] {
        background-color: #1E293B;
        border-right: 1px solid #334155;
    }

    section[data-testid="stSidebar"] h1, 
    section[data-testid="stSidebar"] h2, 
    section[data-testid="stSidebar"] h3,
    section[data-testid="stSidebar"] p,
    section[data-testid="stSidebar"] label,
    section[data-testid="stSidebar"] span,
    section[data-testid="stSidebar"] li,
    section[data-testid="stSidebar"] div[data-testid="stMarkdownContainer"] {
        color: #FFFFFF !important;
    }

    /* Estilo robusto para el expander en la barra lateral */
    section[data-testid="stSidebar"] div[data-testid="stExpander"] {
        background-color: #334155 !important;
        border: 1px solid #475569 !important;
        border-radius: 8px !important;
        margin-top: 10px;
        margin-bottom: 10px;
    }

    section[data-testid="stSidebar"] div[data-testid="stExpander"] details {
        background-color: #334155 !important;
    }

    section[data-testid="stSidebar"] div[data-testid="stExpander"] summary {
        color: #FFFFFF !important;
        font-weight: 600 !important;
    }

    section[data-testid="stSidebar"] div[data-testid="stExpander"] summary:hover {
        color: #FDE047 !important;
    }

    section[data-testid="stSidebar"] div[data-testid="stExpander"] p,
    section[data-testid="stSidebar"] div[data-testid="stExpander"] li {
        color: #F1F5F9 !important;
        font-size: 0.9rem !important;
    }

    section[data-testid="stSidebar"] div[data-testid="stExpander"] a {
        color: #FDE047 !important;
        text-decoration: underline !important;
        font-weight: 600 !important;
    }

    .header-box {
        background-color: #0F172A;
        padding: 22px;
        border-radius: 12px;
        border-left: 6px solid #D4AF37;
        box-shadow: 0 4px 12px rgba(15, 23, 42, 0.2);
        margin-bottom: 24px;
    }

    .header-box h1 {
        color: #FFFFFF !important;
        margin: 0 !important;
        font-size: 1.65rem !important;
        font-weight: 600 !important;
    }

    .header-box p {
        color: #E2E8F0 !important;
        margin-top: 6px !important;
        margin-bottom: 0 !important;
        font-size: 0.95rem !important;
    }

    /* Botón de la barra lateral */
    section[data-testid="stSidebar"] .stButton>button {
        background-color: #FFFFFF !important;
        color: #0F172A !important;
        border: 1px solid #FFFFFF !important;
        border-radius: 6px !important;
        font-weight: 600 !important;
        transition: all 0.3s ease !important;
        width: 100%;
        padding: 8px 12px;
    }

    section[data-testid="stSidebar"] .stButton>button p,
    section[data-testid="stSidebar"] .stButton>button span {
        color: #0F172A !important;
    }

    section[data-testid="stSidebar"] .stButton>button:hover {
        background-color: #0F172A !important;
        color: #FFFFFF !important;
        border-color: #0F172A !important;
    }

    section[data-testid="stSidebar"] .stButton>button:hover p,
    section[data-testid="stSidebar"] .stButton>button:hover span {
        color: #FFFFFF !important;
    }

    .stMainBlockContainer div.stButton > button {
        background-color: #FFFFFF !important;
        color: #0F172A !important;
        border: 1px solid #CBD5E1 !important;
        border-radius: 8px !important;
        padding: 10px 14px !important;
        font-weight: 500 !important;
        width: 100%;
        box-shadow: 0 2px 4px rgba(0,0,0,0.03);
    }

    .stMainBlockContainer div.stButton > button:hover {
        background-color: #0F172A !important;
        color: #FFFFFF !important;
        border-color: #0F172A !important;
    }

    .stTextInput>div>div>input {
        background-color: #FFFFFF !important;
        color: #000000 !important;
        border: 1px solid #CBD5E1 !important;
        border-radius: 6px !important;
    }

    [data-testid="stChatMessage"] {
        background-color: #FFFFFF !important;
        border: 1px solid #E5E7EB !important;
        border-left: 5px solid #0F172A !important;
        border-radius: 10px !important;
        padding: 16px !important;
        margin-bottom: 14px !important;
        box-shadow: 0 2px 8px rgba(0, 0, 0, 0.04) !important;
    }

    [data-testid="stChatMessage"] p, 
    [data-testid="stChatMessage"] div,
    [data-testid="stChatMessage"] span {
        color: #0F172A !important;
        font-size: 0.98rem;
        line-height: 1.6;
    }

    [data-testid="stChatInput"] {
        background-color: #FFFFFF !important;
        border: 2px solid #475569 !important;
        border-radius: 10px !important;
        box-shadow: 0 4px 12px rgba(0, 0, 0, 0.05) !important;
    }

    [data-testid="stChatInput"] textarea {
        color: #000000 !important;
        background-color: #FFFFFF !important;
    }

    [data-testid="stChatInput"] button {
        background-color: #000000 !important;
        border-color: #000000 !important;
        color: #FFFFFF !important;
        border-radius: 6px !important;
    }

    [data-testid="stChatInput"] button svg {
        fill: #FFFFFF !important;
        color: #FFFFFF !important;
    }

    hr {
        border-color: #475569;
    }
    </style>
""",
    unsafe_allow_html=True,
)

# Encabezado superior
st.markdown(
    """
    <div class="header-box">
        <h1>Sistema Integral de Convivencia Escolar</h1>
        <p>Institución Educativa Técnica Sagrado Corazón</p>
    </div>
""",
    unsafe_allow_html=True,
)

api_key = st.secrets.get("GEMINI_API_KEY", "")


def limpiar_texto_para_word(texto: str) -> str:
    if not texto:
        return ""
    texto = re.sub(r"^\s*[\*\-]\s+", "", texto, flags=re.MULTILINE)
    return texto.strip()


def extraer_solo_documento(texto_contenido: str) -> str:
    if not texto_contenido:
        return ""

    patrones = [
        r"(ACTA DE COMPROMISO ESTUDIANTIL.*)",
        r"(ACTA DE COMPROMISO.*)",
        r"(MODELO DE CARTA.*)",
        r"(MODELO DE REGISTRO.*)",
        r"(REGISTRO EN EL OBSERVADOR.*)",
        r"(CARTA DE DESCARGOS.*)",
        r"(CITACIÓN A ACUDIENTE.*)",
        r"(5\.\s*MODELO.*)",
        r"(5\.\s*DOCUMENTO.*)",
        r"(5\.\s*ACTA.*)",
        r"(5\.\s*REGISTRO.*)",
    ]

    for patron in patrones:
        match = re.search(patron, texto_contenido, re.IGNORECASE | re.DOTALL)
        if match:
            texto_extraido = match.group(1)
            lineas = texto_extraido.split("\n")
            if re.match(
                r"^5\.\s*(MODELO|DOCUMENTO|PLANTILLA|ACTA|REGISTRO)",
                lineas[0].strip(),
                re.IGNORECASE,
            ):
                texto_extraido = "\n".join(lineas[1:])
            return texto_extraido.strip()

    return texto_contenido.strip()


def generar_documento_word(texto_contenido):
    doc = Document()
    texto_documento = extraer_solo_documento(texto_contenido)

    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        header = section.header
        p_head = header.paragraphs[0]
        p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p_head.paragraph_format.space_after = Pt(0)

        posibles_nombres_imagen = [
            "escudo.png",
            "escudo.jpg",
            "escudo_intesac.png",
            "escudo_intesac.jpg",
            "logo.png",
            "logo.jpg",
        ]
        imagen_encontrada = None
        for nombre in posibles_nombres_imagen:
            if os.path.exists(nombre):
                imagen_encontrada = nombre
                break

        if imagen_encontrada:
            try:
                r_img = p_head.add_run()
                r_img.add_picture(imagen_encontrada, width=Inches(0.55))
                p_head.add_run("\n")
            except Exception:
                pass

        r_head = p_head.add_run(
            "INSTITUCIÓN EDUCATIVA TÉCNICA SAGRADO CORAZÓN\nSistema Digital de Seguimiento y Convivencia"
        )
        r_head.font.size = Pt(8)
        r_head.font.name = "Arial"
        r_head.font.bold = True
        r_head.font.color.rgb = RGBColor(100, 116, 139)

        footer = section.footer
        p_foot = footer.paragraphs[0]
        p_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_foot = p_foot.add_run(
            "Institución Educativa Técnica Sagrado Corazón — Archivo Digital de Convivencia Escolar"
        )
        r_foot.font.size = Pt(8.5)
        r_foot.font.name = "Arial"
        r_foot.font.italic = True
        r_foot.font.color.rgb = RGBColor(100, 116, 139)

    lineas = texto_documento.split("\n")
    for linea in lineas:
        linea_limpia = limpiar_texto_para_word(linea)
        if not linea_limpia:
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15

        if (
            "ACTA" in linea.upper()
            or "REGISTRO" in linea.upper()
            or "MODELO" in linea.upper()
            or linea.strip().startswith("#")
        ):
            texto_titulo = linea_limpia.replace("#", "").replace("**", "")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(texto_titulo)
            run.bold = True
            run.font.size = Pt(12)
            run.font.name = "Arial"
            run.font.color.rgb = RGBColor(15, 23, 42)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(12)

        elif "________________" in linea or "FIRMA" in linea.upper() or "T.I." in linea.upper() or "C.C." in linea.upper():
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(linea_limpia.replace("**", ""))
            run.font.size = Pt(10)
            run.font.name = "Arial"
            run.font.bold = "FIRMA" in linea.upper() or "ESTUDIANTE" in linea.upper()
            run.font.color.rgb = RGBColor(31, 41, 55)
            if "________________" in linea and "Firma" in linea:
                p.paragraph_format.space_before = Pt(20)
                p.paragraph_format.space_after = Pt(2)
            else:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(4)

        else:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            run = p.add_run(linea_limpia.replace("**", ""))
            run.font.size = Pt(10)
            run.font.name = "Arial"
            run.font.color.rgb = RGBColor(31, 41, 55)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


# Barra lateral izquierda
with st.sidebar:
    st.header("Configuración del Sistema")

    if not api_key:
        api_key = st.text_input("Clave API de Gemini", type="password")

    user_role = st.selectbox(
        "Perfil del Consultante:",
        ["Estudiante", "Acudiente / Padre de Familia", "Docente / Directivo"],
    )

    st.markdown("---")
    st.markdown("### Biblioteca Virtual")
    st.markdown(
        "[Abrir Repositorio Digital](https://repositorioconvivencia-k3hz5bcgykhxqwn6gn7cjh.streamlit.app/)"
    )
    
    with st.expander("Ver Marco Legal e Institucional"):
        st.markdown(
            f"* **[Constitución Política]({URL_CONSTITUCION_POLITICA})**: Art. 29 (Garantía del Debido Proceso)."
        )
        st.markdown(
            f"* **[Ley 115 de 1994]({URL_LEY_115})**: Ley General de Educación."
        )
        st.markdown(
            f"* **[Ley 1098 de 2006]({URL_LEY_1098})**: Código de la Infancia y la Adolescencia."
        )
        st.markdown(
            f"* **[Ley 1620 de 2013]({URL_LEY_1620})**: Sistema Nacional de Convivencia Escolar."
        )
        st.markdown(
            f"* **[Decreto 1965 de 2013]({URL_DECRETO_1965})**: Decreto Reglamentario de la Ley 1620."
        )
        st.markdown(
            f"* **[Manual de Convivencia]({URL_MANUAL_CONVIVENCIA})**: Manual Institucional."
        )

    st.markdown("---")
    st.markdown("**Guía de consulta:**")
    st.markdown("1. Ingrese los detalles de la situación.")
    st.markdown("2. Responda a la solicitud de datos o elija plantilla en blanco.")
    st.markdown("3. Descargue el documento oficial en Word generado al final.")

    st.markdown("---")
    if st.button("Reiniciar consulta", use_container_width=True):
        st.session_state.messages = []
        st.rerun()

SYSTEM_PROMPT = f"""
Eres el asistente institucional del Sistema Digital de Llamados de Atención y Seguimiento de Convivencia Escolar de la Institución Educativa Técnica Sagrado Corazón de Soledad.

Estás orientando a un usuario con el perfil de: {user_role}.

Tu propósito es asesorar formal y pedagógicamente a la comunidad educativa ante situaciones disciplinarias, asegurando el cumplimiento de la Constitución Política de Colombia (Art. 29 - Debido Proceso y el derecho de defensa del estudiante), la Ley 115 de 1994, la Ley 1098 de 2006 (Código de Infancia y Adolescencia), la Ley 1620 de 2013, el Decreto 1965 de 2013 y el Manual de Convivencia de la Institución Educativa Técnica Sagrado Corazón.

REGLAS ESTRICTAS DE INTERACCIÓN Y FLUJO:
1. **Fase 1 (Análisis Inicial, Derechos de Defensa y Solicitud de Datos - NUNCA GENERES EL DOCUMENTO AQUÍ)**: 
   Cuando el usuario ingrese o seleccione un caso, preséntale ÚNICAMENTE:
   - Resumen breve de la situación reportada.
   - Clasificación de la falta (Tipo I, II o III según Manual de Convivencia y Ley 1620).
   - Procedimiento institucional, **los derechos del estudiante y sus mecanismos de defensa** (derecho a ser escuchado, presentar descargos, aportar pruebas y contradecir en el marco del Debido Proceso del Artículo 29 de la Constitución Política).
   - **Pregunta obligatoria al final**: Pregúntale claramente al usuario: "¿Desea que genere el Acta de Compromiso completando automáticamente los datos de este caso (por favor proporcione: nombre completo del estudiante, curso, número de T.I. - Tarjeta de Identidad y la fecha del hecho o documento), o prefiere una plantilla en blanco con líneas de subrayado (`____________________`) para diligenciarla manualmente?"
   *¡IMPORTANTE!* En este primer mensaje **NO** debes redactar ni incluir ningún formato de acta ni documento. Solo analiza y haz la pregunta.

2. **Fase 2 (Generación del Documento - Únicamente al recibir respuesta del usuario)**: 
   Solo cuando el usuario indique su preferencia y proporcione los datos (o pida la plantilla en blanco), redactarás y generarás el documento oficial cumpliendo **estrictamente** esta estructura visual y formal basada en los formatos institucionales:

   - Título centrado: 
     ACTA DE COMPROMISO ESTUDIANTIL

   - Fecha estructurada: 
     Soledad, [Fecha proporcionada o líneas de subrayado].

   - Párrafo introductorio formal:
     Yo, [Nombre del estudiante o línea de subrayado], identificado(a) con T.I. N° [Número de T.I. o línea], estudiante del curso [Curso], en ejercicio de mis derechos y del debido proceso, prometo solemnemente tener una conducta correcta, responsable, respetuosa, acudir puntualmente a mis horas de clase y cumplir con todas las actividades académicas y de convivencia que me corresponden como estudiante de la Institución Educativa Técnica Sagrado Corazón.

   - Compromisos específicos (derivados del caso analizado en lista numerada 1., 2., 3.).

   - Cláusula de incumplimiento:
     En caso de no cumplir con los compromisos establecidos, se me aplicarán las sanciones y correctivos correspondientes conforme al Manual de Convivencia institucional, garantizando en todo momento el derecho de defensa.

   - Bloque de Firmas profesional al pie:
     _____________________________________
     Firma del Estudiante
     T.I. ________________________________

     _____________________________________
     Firma del Acudiente / Representante
     C.C. ________________________________

     _____________________________________
     Docente Tutor / Coordinador(a)
"""

WELCOME_MESSAGE = f"""
Saludos. Bienvenido(a) al Sistema Digital de Llamados de Atención y Seguimiento de Convivencia Escolar de la Institución Educativa Técnica Sagrado Corazón.

Sesión iniciada como: **{user_role}**.

Por favor, seleccione una de las situaciones predeterminadas a continuación o redacte detalladamente lo sucedido en la casilla de texto inferior.
"""

if "messages" not in st.session_state or len(st.session_state.messages) == 0:
    st.session_state.messages = [
        {"role": "assistant", "content": WELCOME_MESSAGE}
    ]

for idx, msg in enumerate(st.session_state.messages):
    with st.chat_message(msg["role"]):
        st.markdown(msg["content"])

        contenido_upper = msg["content"].upper()
        # Condición estricta: Solo muestra el botón de descarga si el mensaje contiene el bloque de firmas o la estructura final del documento
        if (
            msg["role"] == "assistant"
            and idx > 0
            and (
                "FIRMA DEL ESTUDIANTE" in contenido_upper
                or "_____________________________________" in msg["content"]
            )
        ):
            if HAS_DOCX:
                docx_bytes = generar_documento_word(msg["content"])
                st.download_button(
                    label="📄 Descargar Acta Oficial en Microsoft Word (.docx)",
                    data=docx_bytes,
                    file_name=f"Acta_Compromiso_SagradoCorazon_{idx}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    key=f"dl_word_{idx}",
                )
            else:
                st.download_button(
                    label="📄 Guardar texto (.txt)",
                    data=msg["content"],
                    file_name=f"Acta_Compromiso_SagradoCorazon_{idx}.txt",
                    mime="text/plain",
                    key=f"dl_txt_{idx}",
                )

selected_option = None
if len(st.session_state.messages) <= 1:
    st.markdown("**Seleccione el tipo de situación o escriba su caso abajo:**")
    col1, col2 = st.columns(2)
    with col1:
        if st.button("Exceso de maquillaje"):
            selected_option = "Se reporta un llamado de atención por exceso de maquillaje o incumplimiento del código de presentación personal."
        if st.button("Conflictos / Agresión verbal o física"):
            selected_option = "Ocurrió una situación de conflicto o agresión entre estudiantes dentro de la institución educativa."
        if st.button("Presunto Acoso Escolar (Bullying)"):
            selected_option = "Se presenta una situación reiterada de presunto acoso escolar (bullying) o ciberacoso."
        if st.button("Fraude académico / Plagio"):
            selected_option = "Se reporta una falta relacionada con fraude académico o plagio en evaluación durante la clase."
        if st.button("Desacato o falta de respeto a docente"):
            selected_option = "Se presentó un acto de desobediencia o falta de respeto verbal hacia un docente durante el desarrollo de la clase."

    with col2:
        if st.button("Corte de cabello / Uniforme"):
            selected_option = "Se presenta un llamado de atención por corte de cabello inadecuado o porte incorrecto del uniforme institucional."
        if st.button("Incumplimiento de deberes / Asistencia"):
            selected_option = "Se presentó un incumplimiento en los deberes académicos, faltas de asistencia o impuntualidad."
        if st.button("Uso no autorizado de celular/equipos"):
            selected_option = "Se reporta el uso no autorizado de teléfono celular o dispositivos electrónicos durante la jornada escolar."
        if st.button("Evasión de clase / Ausencia en aula"):
            selected_option = "El estudiante ingresó a la institución pero evadió la clase o se ausentó del aula."
        if st.button("Daño a propiedad institucional"):
            selected_option = "Se reportan daños materiales a los pupitres, paredes u otros bienes de la institución."

user_input = st.chat_input("Escriba aquí los hechos de la situación o los datos solicitados...")
prompt = selected_option or user_input

if prompt:
    if not api_key:
        st.warning("Se requiere ingresar la Clave API para continuar.")
        st.stop()

    st.session_state.messages.append({"role": "user", "content": prompt})

    try:
        client = genai.Client(api_key=api_key)

        contents = []
        for m in st.session_state.messages:
            role = "user" if m["role"] == "user" else "model"
            contents.append(
                types.Content(
                    role=role,
                    parts=[types.Part.from_text(text=m["content"])],
                )
            )

        with st.spinner("Procesando información institucional..."):
            response = client.models.generate_content(
                model="gemini-3.6-flash",
                contents=contents,
                config=types.GenerateContentConfig(
                    system_instruction=SYSTEM_PROMPT, temperature=0.2
                ),
            )

            if response and response.text:
                st.session_state.messages.append(
                    {"role": "assistant", "content": response.text}
                )
                st.rerun()

    except Exception as e:
        error_msg = str(e)
        if "429" in error_msg or "RESOURCE_EXHAUSTED" in error_msg:
            st.warning(
                "⚠️ El servicio ha alcanzado el límite de consultas por minuto de la capa gratuita. Por favor, espere 30 a 40 segundos e intente de nuevo."
            )
        else:
            st.error(f"Error de comunicación con el servicio: {e}")
